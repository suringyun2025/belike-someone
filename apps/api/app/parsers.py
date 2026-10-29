from __future__ import annotations

import csv
import io
import json
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import settings
from .security import redact_text, sha256_text


class ParseError(Exception):
    pass


class NeedsOCR(ParseError):
    pass


class NeedsASR(ParseError):
    pass


@dataclass
class ParsedBlock:
    text: str
    kind: str = "paragraph"
    heading: str = ""
    page: int | None = None
    start_char: int = 0
    end_char: int = 0
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    blocks: list[ParsedBlock]
    parser: str
    parser_version: str = "1"
    language: str = "zh"
    meta: dict[str, Any] = field(default_factory=dict)
    redaction_count: int = 0

    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks if block.text.strip())


def detect_language(text: str) -> str:
    sample = text[:5000]
    cjk = len(re.findall(r"[\u4e00-\u9fff]", sample))
    latin = len(re.findall(r"[A-Za-z]", sample))
    return "zh" if cjk >= max(10, latin * 0.2) else "en"


def estimate_tokens(text: str) -> int:
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    words = len(re.findall(r"[A-Za-z0-9_]+", text))
    return max(1, int(cjk * 0.8 + words * 1.2))


def _split_paragraphs(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    parts = re.split(r"\n\s*\n", normalized)
    return [part.strip() for part in parts if part.strip()]


def _finalize_blocks(blocks: list[ParsedBlock]) -> ParsedDocument:
    offset = 0
    clean: list[ParsedBlock] = []
    for block in blocks:
        text = block.text.strip()
        if not text:
            continue
        text, count = redact_text(text)
        block.text = text
        block.start_char = offset
        block.end_char = offset + len(text)
        offset = block.end_char + 2
        clean.append(block)
    combined = "\n\n".join(block.text for block in clean)
    return ParsedDocument(
        blocks=clean,
        parser="builtin",
        language=detect_language(combined),
        redaction_count=sum(
            1 for block in clean for _ in re.finditer(r"\[REDACTED:[A-Z_]+\]", block.text)
        ),
    )


def _parse_markdown(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[ParsedBlock] = []
    heading_stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_buffer() -> None:
        nonlocal buffer
        if buffer:
            heading = " / ".join(item[1] for item in heading_stack)
            blocks.append(ParsedBlock("\n".join(buffer).strip(), heading=heading))
            buffer = []

    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(
                    ParsedBlock(
                        "\n".join(code_lines).strip(),
                        kind="code",
                        heading=" / ".join(item[1] for item in heading_stack),
                    )
                )
                code_lines = []
                in_code = False
            else:
                flush_buffer()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            flush_buffer()
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            heading_stack = [item for item in heading_stack if item[0] < level]
            heading_stack.append((level, title))
            blocks.append(
                ParsedBlock(
                    title,
                    kind="heading",
                    heading=" / ".join(item[1] for item in heading_stack),
                    meta={"level": level},
                )
            )
            continue
        if not line.strip():
            flush_buffer()
        else:
            buffer.append(line)
    flush_buffer()
    if in_code and code_lines:
        blocks.append(ParsedBlock("\n".join(code_lines).strip(), kind="code"))
    document = _finalize_blocks(blocks)
    document.parser = "markdown"
    return document


def _parse_plain_text(path: Path, kind: str = "paragraph") -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    blocks = [ParsedBlock(part, kind=kind) for part in _split_paragraphs(text)]
    document = _finalize_blocks(blocks)
    document.parser = "text"
    return document


_OCR_ENGINE = None


def _get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError as exc:
            raise NeedsOCR(
                "扫描 PDF / 图片需要 OCR 扩展；请执行 uv sync --project apps/api --extra ocr 后重试"
            ) from exc
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_path(path: Path) -> tuple[str, float]:
    engine = _get_ocr_engine()
    result, _elapsed = engine(str(path))
    texts: list[str] = []
    scores: list[float] = []
    for item in result or []:
        if len(item) >= 3 and str(item[1]).strip():
            texts.append(str(item[1]).strip())
            try:
                scores.append(float(item[2]))
            except (TypeError, ValueError):
                pass
    text = "\n".join(texts)
    if not text:
        raise NeedsOCR("OCR 未识别到可用文字")
    return text, (sum(scores) / len(scores) if scores else 0.0)


def _parse_image(path: Path) -> ParsedDocument:
    text, confidence = _ocr_path(path)
    blocks = [
        ParsedBlock(
            text,
            kind="ocr",
            page=1,
            meta={"ocr": True, "confidence": round(confidence, 4)},
        )
    ]
    document = _finalize_blocks(blocks)
    document.parser = "rapidocr"
    document.meta = {"ocr": True, "confidence": round(confidence, 4)}
    return document


def _parse_pdf_ocr(path: Path, page_count: int) -> ParsedDocument:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise NeedsOCR(
            "扫描 PDF 需要 pypdfium2 与 OCR 扩展；请执行 uv sync --project apps/api --extra ocr"
        ) from exc
    pdf = pdfium.PdfDocument(str(path))
    blocks: list[ParsedBlock] = []
    confidences: list[float] = []
    try:
        with tempfile.TemporaryDirectory(prefix="distill-ocr-") as tmp:
            for index in range(page_count):
                page = pdf[index]
                image = page.render(scale=2).to_pil()
                image_path = Path(tmp) / f"page-{index + 1}.png"
                image.save(image_path)
                text, confidence = _ocr_path(image_path)
                confidences.append(confidence)
                blocks.append(
                    ParsedBlock(
                        text,
                        kind="ocr",
                        page=index + 1,
                        heading=f"第 {index + 1} 页",
                        meta={"ocr": True, "confidence": round(confidence, 4)},
                    )
                )
                page.close()
    finally:
        pdf.close()
    if not blocks:
        raise NeedsOCR("PDF OCR 未识别到可用文字")
    document = _finalize_blocks(blocks)
    document.parser = "pypdfium2+rapidocr"
    document.meta = {
        "ocr": True,
        "pages": page_count,
        "confidence": round(sum(confidences) / len(confidences), 4) if confidences else 0,
    }
    return document


def _parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    blocks: list[ParsedBlock] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if not page_text:
            continue
        for part in _split_paragraphs(page_text):
            blocks.append(ParsedBlock(part, page=index, heading=f"第 {index} 页"))
    if not blocks:
        return _parse_pdf_ocr(path, len(reader.pages))
    document = _finalize_blocks(blocks)
    document.parser = "pypdf"
    document.meta = {"pages": len(reader.pages)}
    return document


def _parse_media(path: Path) -> ParsedDocument:
    try:
        from faster_whisper import WhisperModel
    except ImportError as exc:
        raise NeedsASR(
            "音频/视频转录需要 ASR 扩展；请执行 uv sync --project apps/api --extra asr 后重试"
        ) from exc
    local_model = settings.data_dir / "models" / "faster-whisper-base"
    model_name = str(local_model) if local_model.exists() else settings.whisper_model
    device = settings.whisper_device
    model = WhisperModel(model_name, device=device, compute_type="int8")
    segments, info = model.transcribe(str(path), vad_filter=True)
    blocks: list[ParsedBlock] = []
    for segment in segments:
        text = (segment.text or "").strip()
        if not text:
            continue
        blocks.append(
            ParsedBlock(
                text,
                kind="transcript",
                meta={
                    "start": round(segment.start, 2),
                    "end": round(segment.end, 2),
                },
            )
        )
    if not blocks:
        raise NeedsASR("ASR 未识别到可用语音文本")
    document = _finalize_blocks(blocks)
    document.parser = "faster-whisper"
    document.meta = {
        "language": getattr(info, "language", ""),
        "duration": getattr(info, "duration", None),
        "model": model_name,
    }
    return document


def _parse_docx(path: Path) -> ParsedDocument:
    doc = DocxDocument(str(path))
    blocks: list[ParsedBlock] = []
    heading_stack: list[tuple[int, str]] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        level_match = re.search(r"(\d+)", style)
        if style.lower().startswith("heading") and level_match:
            level = int(level_match.group(1))
            heading_stack = [item for item in heading_stack if item[0] < level]
            heading_stack.append((level, text))
            blocks.append(
                ParsedBlock(
                    text,
                    kind="heading",
                    heading=" / ".join(item[1] for item in heading_stack),
                )
            )
        else:
            blocks.append(
                ParsedBlock(text, heading=" / ".join(item[1] for item in heading_stack))
            )
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(
                    ParsedBlock(
                        " | ".join(cells),
                        kind="table",
                        heading=" / ".join(item[1] for item in heading_stack),
                    )
                )
    document = _finalize_blocks(blocks)
    document.parser = "python-docx"
    return document


def _parse_html(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    blocks: list[ParsedBlock] = []
    heading_stack: list[tuple[int, str]] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "td"]):
        text = element.get_text(" ", strip=True)
        if not text:
            continue
        if element.name.startswith("h"):
            level = int(element.name[1])
            heading_stack = [item for item in heading_stack if item[0] < level]
            heading_stack.append((level, text))
            blocks.append(
                ParsedBlock(
                    text,
                    kind="heading",
                    heading=" / ".join(item[1] for item in heading_stack),
                )
            )
        else:
            kind = "code" if element.name == "pre" else "paragraph"
            blocks.append(
                ParsedBlock(
                    text,
                    kind=kind,
                    heading=" / ".join(item[1] for item in heading_stack),
                )
            )
    document = _finalize_blocks(blocks)
    document.parser = "beautifulsoup"
    return document


def _parse_xlsx(path: Path) -> ParsedDocument:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[ParsedBlock] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        start_row = 1
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if not values:
                continue
            rows.append(" | ".join(values))
            if len(rows) >= 12:
                blocks.append(
                    ParsedBlock(
                        "\n".join(rows),
                        kind="table",
                        heading=f"{sheet.title} 第 {start_row}-{index} 行",
                        meta={"sheet": sheet.title, "row_start": start_row, "row_end": index},
                    )
                )
                rows = []
                start_row = index + 1
        if rows:
            blocks.append(
                ParsedBlock(
                    "\n".join(rows),
                    kind="table",
                    heading=f"{sheet.title} 第 {start_row} 行起",
                    meta={"sheet": sheet.title, "row_start": start_row},
                )
            )
    document = _finalize_blocks(blocks)
    document.parser = "openpyxl"
    return document


def _parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    rows: list[str] = []
    blocks: list[ParsedBlock] = []
    header: list[str] = []
    for index, row in enumerate(reader, start=1):
        if not row or not any(cell.strip() for cell in row):
            continue
        if not header:
            header = [cell.strip() for cell in row]
        rows.append(" | ".join(cell.strip() for cell in row))
        if len(rows) >= 15:
            blocks.append(
                ParsedBlock(
                    "\n".join(rows),
                    kind="table",
                    heading=f"第 {max(1, index - 14)}-{index} 行",
                    meta={"row_end": index},
                )
            )
            rows = []
    if rows:
        blocks.append(ParsedBlock("\n".join(rows), kind="table"))
    document = _finalize_blocks(blocks)
    document.parser = "csv"
    document.meta = {"columns": header}
    return document


def _looks_like_message(item: Any) -> bool:
    return isinstance(item, dict) and any(
        key in item for key in ("content", "text", "message", "msg")
    )


def _message_text(item: dict[str, Any]) -> str:
    for key in ("content", "text", "message", "msg"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _message_sender(item: dict[str, Any]) -> str:
    for key in ("sender", "from", "author", "name", "user", "role"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _parse_json_messages(path: Path, jsonl: bool = False) -> ParsedDocument:
    items: list[Any] = []
    if jsonl:
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                items.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    else:
        payload = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        if isinstance(payload, dict) and isinstance(payload.get("messages"), list):
            items = payload["messages"]
        elif isinstance(payload, list):
            items = payload
        elif _looks_like_message(payload):
            items = [payload]
        else:
            items = [payload]

    blocks: list[ParsedBlock] = []
    if items and all(_looks_like_message(item) for item in items):
        for index, item in enumerate(items):
            text = _message_text(item)
            if not text:
                continue
            timestamp = item.get("time") or item.get("timestamp") or item.get("date") or ""
            sender = _message_sender(item)
            prefix = f"[{timestamp}] {sender}: " if timestamp or sender else ""
            blocks.append(
                ParsedBlock(
                    prefix + text,
                    kind="chat",
                    meta={
                        "message_index": index,
                        "sender": sender,
                        "timestamp": str(timestamp),
                    },
                )
            )
        parser = "jsonl-chat" if jsonl else "json-chat"
    else:
        text = json.dumps(items, ensure_ascii=False, indent=2)
        blocks = [ParsedBlock(part) for part in _split_paragraphs(text)]
        parser = "json"
    document = _finalize_blocks(blocks)
    document.parser = parser
    return document


_CHAT_LINE_RE = re.compile(
    r"^\s*(?:\[(?P<time>[^\]]+)\]\s*)?(?P<sender>[^:：]{1,40})[:：]\s*(?P<text>.+)$"
)


def _parse_chat_text(path: Path) -> ParsedDocument:
    blocks: list[ParsedBlock] = []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        match = _CHAT_LINE_RE.match(line)
        if match:
            blocks.append(
                ParsedBlock(
                    line.strip(),
                    kind="chat",
                    meta={
                        "sender": match.group("sender").strip(),
                        "timestamp": (match.group("time") or "").strip(),
                    },
                )
            )
        elif line.strip():
            blocks.append(ParsedBlock(line.strip(), kind="chat"))
    document = _finalize_blocks(blocks)
    document.parser = "chat-text"
    return document


def _parse_code(path: Path) -> ParsedDocument:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    blocks: list[ParsedBlock] = []
    for start in range(0, len(lines), 80):
        part = "\n".join(lines[start : start + 80]).strip()
        if not part:
            continue
        blocks.append(
            ParsedBlock(
                part,
                kind="code",
                heading=f"{path.name} 第 {start + 1}-{min(start + 80, len(lines))} 行",
                meta={"path": path.name, "line_start": start + 1},
            )
        )
    document = _finalize_blocks(blocks)
    document.parser = "code"
    return document


def parse_file(path: str | Path, filename: str, mime: str = "", source_type: str = "document") -> ParsedDocument:
    file_path = Path(path)
    ext = Path(filename).suffix.lower()
    try:
        if source_type == "media":
            return _parse_media(file_path)
        if source_type == "image":
            return _parse_image(file_path)
        if ext in {".md", ".markdown"}:
            return _parse_markdown(file_path)
        if ext == ".pdf":
            return _parse_pdf(file_path)
        if ext == ".docx":
            return _parse_docx(file_path)
        if ext in {".html", ".htm"}:
            return _parse_html(file_path)
        if ext == ".xlsx":
            return _parse_xlsx(file_path)
        if ext == ".csv":
            return _parse_csv(file_path)
        if ext == ".jsonl":
            return _parse_json_messages(file_path, jsonl=True)
        if ext == ".json":
            return _parse_json_messages(file_path, jsonl=False)
        if source_type == "chat":
            return _parse_chat_text(file_path)
        if source_type == "code":
            return _parse_code(file_path)
        if ext in {".txt", ".log", ".rst", ".yaml", ".yml", ".toml", ".xml"}:
            return _parse_plain_text(file_path)
        raise ParseError(f"暂不支持的文件格式：{ext or filename}")
    except (ParseError, NeedsOCR, NeedsASR):
        raise
    except Exception as exc:  # pragma: no cover - defensive boundary for third-party parsers
        raise ParseError(f"解析失败：{exc}") from exc


def chunk_text(
    blocks: list[ParsedBlock],
    chunk_chars: int = 1400,
    overlap: int = 180,
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    current: list[ParsedBlock] = []
    current_len = 0

    def flush(carry: bool = True) -> None:
        nonlocal current, current_len
        if not current:
            return
        text = "\n\n".join(block.text for block in current)
        pages = [block.page for block in current if block.page is not None]
        headings = [block.heading for block in current if block.heading]
        chunks.append(
            {
                "text": text,
                "heading_path": headings[-1] if headings else "",
                "page_start": min(pages) if pages else None,
                "page_end": max(pages) if pages else None,
                "start_char": current[0].start_char,
                "end_char": current[-1].end_char,
                "token_estimate": estimate_tokens(text),
                "sha256": sha256_text(text),
                "meta": {"kinds": sorted({block.kind for block in current})},
            }
        )
        carry_block = None
        if carry and current and len(current[-1].text) <= overlap:
            carry_block = current[-1]
        current = [carry_block] if carry_block else []
        current_len = len(carry_block.text) if carry_block else 0

    for block in blocks:
        block_len = len(block.text) + 2
        if block.kind == "heading" and current:
            flush(carry=False)
        if current and current_len + block_len > chunk_chars:
            flush()
        current.append(block)
        current_len += block_len
    flush()
    return chunks
